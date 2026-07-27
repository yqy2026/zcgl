"""Deterministic parser-backed validation for staged uploads."""

from __future__ import annotations

import hashlib
import json
import warnings
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Literal

import fitz
from defusedxml import ElementTree as DefusedElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pydantic import BaseModel, ConfigDict, ValidationError

from src.core.exception_handler import BusinessValidationError

from .profiles import UploadFormat, UploadProfile, UploadPurpose, get_upload_profile

_DANGEROUS_SUFFIXES = {
    ".bat",
    ".cmd",
    ".com",
    ".dll",
    ".exe",
    ".gif",
    ".htm",
    ".html",
    ".jar",
    ".jpeg",
    ".jpg",
    ".js",
    ".json",
    ".pdf",
    ".php",
    ".png",
    ".ps1",
    ".rar",
    ".tif",
    ".tiff",
    ".vbs",
    ".xls",
    ".xlsm",
    ".xlsx",
    ".zip",
}
_REQUIRED_XLSX_MEMBERS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
}
_INVALID_FILENAME_CHARACTERS = frozenset('<>:"|?*')


class UploadValidationError(BusinessValidationError):
    def __init__(
        self,
        error_code: str,
        purpose: UploadPurpose,
        message: str,
        **details: object,
    ) -> None:
        super().__init__(
            message,
            details={
                "upload_error_code": error_code,
                "profile": purpose.value,
                **details,
            },
        )


class PasswordPolicyDocument(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    min_length: int
    require_uppercase: bool
    require_lowercase: bool
    require_numbers: bool
    require_special_chars: bool


class SystemSettingsDocument(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    site_name: str
    site_description: str
    allow_registration: bool
    session_timeout: int
    password_policy: PasswordPolicyDocument


class SystemSettingsRestoreDocument(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    backup_time: datetime
    backup_user: str
    system_settings: SystemSettingsDocument
    backup_type: Literal["full"]
    version: Literal["2.0.0"]


@dataclass(frozen=True)
class ValidatedFile:
    purpose: UploadPurpose
    original_filename: str
    canonical_extension: str
    canonical_mime: str
    size_bytes: int
    sha256: str


def validate_upload_filename(filename: str, purpose: UploadPurpose) -> None:
    if (
        filename == ""
        or len(filename) > 255
        or "/" in filename
        or "\\" in filename
        or Path(filename).name != filename
        or any(ord(character) < 32 for character in filename)
        or any(character in _INVALID_FILENAME_CHARACTERS for character in filename)
    ):
        raise UploadValidationError(
            "filename_invalid", purpose, "upload filename is invalid"
        )
    suffixes = [suffix.lower() for suffix in Path(filename).suffixes]
    if any(suffix in _DANGEROUS_SUFFIXES for suffix in suffixes[:-1]):
        raise UploadValidationError(
            "filename_invalid",
            purpose,
            "upload filename has an unsafe double extension",
        )


def _select_format(
    profile: UploadProfile, filename: str, declared_mime: str | None
) -> tuple[str, UploadFormat]:
    extension = Path(filename).suffix.lower()
    file_format = profile.formats.get(extension)
    if file_format is None:
        raise UploadValidationError(
            "extension_not_allowed",
            profile.purpose,
            "upload extension is not allowed for this purpose",
            extension=extension,
        )
    if declared_mime != file_format.canonical_mime:
        raise UploadValidationError(
            "mime_mismatch",
            profile.purpose,
            "declared MIME does not match the upload extension",
            extension=extension,
            declared_mime=declared_mime,
            expected_mime=file_format.canonical_mime,
        )
    return extension, file_format


def _validate_signature(
    path: Path, profile: UploadProfile, file_format: UploadFormat
) -> None:
    with path.open("rb") as source:
        header = source.read(16)
    if file_format.parser == "json":
        header = header.lstrip(b" \t\r\n")
    if file_format.parser == "xlsx" and header.startswith(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    ):
        raise UploadValidationError(
            "xlsx_encrypted", profile.purpose, "encrypted Excel files are not allowed"
        )
    if not any(header.startswith(signature) for signature in file_format.signatures):
        raise UploadValidationError(
            "signature_mismatch",
            profile.purpose,
            "file signature does not match its declared type",
        )


def _validate_pdf(path: Path, profile: UploadProfile) -> None:
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise UploadValidationError(
            "pdf_invalid", profile.purpose, "PDF structure is invalid"
        ) from exc
    try:
        if document.needs_pass or document.is_encrypted:
            raise UploadValidationError(
                "pdf_encrypted", profile.purpose, "encrypted PDFs are not allowed"
            )
        if document.is_repaired:
            raise UploadValidationError(
                "pdf_repaired", profile.purpose, "repaired PDFs are not allowed"
            )
        if document.page_count == 0:
            raise UploadValidationError(
                "pdf_zero_pages", profile.purpose, "PDF must contain at least one page"
            )
        if (
            profile.max_pdf_pages is not None
            and document.page_count > profile.max_pdf_pages
        ):
            raise UploadValidationError(
                "pdf_page_limit",
                profile.purpose,
                "PDF page limit exceeded",
                page_count=document.page_count,
                max_pages=profile.max_pdf_pages,
            )
        for page_number in range(document.page_count):
            document.load_page(page_number)
    except UploadValidationError:
        raise
    except Exception as exc:
        raise UploadValidationError(
            "pdf_invalid", profile.purpose, "PDF structure is invalid"
        ) from exc
    finally:
        document.close()


def _validate_image(path: Path, profile: UploadProfile, extension: str) -> None:
    expected_format = "PNG" if extension == ".png" else "JPEG"
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                width, height = image.size
                if width <= 0 or height <= 0:
                    raise ValueError("image dimensions must be positive")
                pixels = width * height
                if (
                    profile.max_image_pixels is not None
                    and pixels > profile.max_image_pixels
                ):
                    raise UploadValidationError(
                        "image_pixel_limit",
                        profile.purpose,
                        "image pixel limit exceeded",
                        pixels=pixels,
                        max_pixels=profile.max_image_pixels,
                    )
                if image.format != expected_format:
                    raise UploadValidationError(
                        "image_format_mismatch",
                        profile.purpose,
                        "decoded image format does not match its extension",
                    )
                image.verify()
            with Image.open(path) as image:
                image.load()
    except UploadValidationError:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise UploadValidationError(
            "image_pixel_limit", profile.purpose, "image pixel limit exceeded"
        ) from exc
    except (OSError, SyntaxError, ValueError, UnidentifiedImageError) as exc:
        raise UploadValidationError(
            "image_invalid", profile.purpose, "image structure is invalid"
        ) from exc


def _validate_xlsx(path: Path, profile: UploadProfile) -> None:
    limits = profile.xlsx_limits
    if limits is None:
        raise RuntimeError("XLSX profile is missing resource limits")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > limits.max_members:
                raise UploadValidationError(
                    "xlsx_resource_limit",
                    profile.purpose,
                    "XLSX member limit exceeded",
                )
            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise UploadValidationError(
                    "xlsx_duplicate_member",
                    profile.purpose,
                    "XLSX contains duplicate members",
                )
            total_uncompressed = 0
            for member in members:
                normalized = member.filename.replace("\\", "/")
                member_path = PurePosixPath(normalized)
                if (
                    normalized.startswith("/")
                    or member_path.is_absolute()
                    or ".." in member_path.parts
                ):
                    raise UploadValidationError(
                        "xlsx_path_unsafe",
                        profile.purpose,
                        "XLSX contains an unsafe member path",
                    )
                if member.flag_bits & 0x1:
                    raise UploadValidationError(
                        "xlsx_encrypted",
                        profile.purpose,
                        "encrypted XLSX members are not allowed",
                    )
                if normalized.casefold().endswith("vbaproject.bin"):
                    raise UploadValidationError(
                        "xlsx_macro_forbidden",
                        profile.purpose,
                        "macro-enabled workbooks are not allowed",
                    )
                if member.file_size > limits.max_member_bytes:
                    raise UploadValidationError(
                        "xlsx_resource_limit",
                        profile.purpose,
                        "XLSX member size limit exceeded",
                    )
                total_uncompressed += member.file_size
                if total_uncompressed > limits.max_total_uncompressed_bytes:
                    raise UploadValidationError(
                        "xlsx_resource_limit",
                        profile.purpose,
                        "XLSX total uncompressed size limit exceeded",
                    )
                if member.file_size > 0 and (
                    member.compress_size == 0
                    or member.file_size / member.compress_size
                    > limits.max_compression_ratio
                ):
                    raise UploadValidationError(
                        "xlsx_resource_limit",
                        profile.purpose,
                        "XLSX compression ratio limit exceeded",
                    )
            missing = _REQUIRED_XLSX_MEMBERS.difference(names)
            if missing:
                raise UploadValidationError(
                    "xlsx_required_member_missing",
                    profile.purpose,
                    "XLSX is missing required OOXML members",
                    missing_members=sorted(missing),
                )
            for member in members:
                if member.filename.endswith((".xml", ".rels")):
                    try:
                        with archive.open(member) as xml_stream:
                            tree = DefusedElementTree.parse(xml_stream)
                        if member.filename == "[Content_Types].xml" and any(
                            "macroenabled" in str(value).casefold()
                            or "vbaproject" in str(value).casefold()
                            for element in tree.iter()
                            for value in element.attrib.values()
                        ):
                            raise UploadValidationError(
                                "xlsx_macro_forbidden",
                                profile.purpose,
                                "macro-enabled workbooks are not allowed",
                            )
                    except (DefusedXmlException, DefusedElementTree.ParseError) as exc:
                        raise UploadValidationError(
                            "xlsx_xml_unsafe",
                            profile.purpose,
                            "XLSX contains unsafe or invalid XML",
                        ) from exc
    except UploadValidationError:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise UploadValidationError(
            "xlsx_invalid", profile.purpose, "XLSX archive is invalid"
        ) from exc

    try:
        with path.open("rb") as source:
            workbook = load_workbook(
                source, read_only=True, data_only=False, keep_vba=False
            )
            workbook.close()
    except Exception as exc:
        raise UploadValidationError(
            "xlsx_invalid", profile.purpose, "XLSX workbook is invalid"
        ) from exc


def _validate_json(path: Path, profile: UploadProfile) -> None:
    try:
        text = path.read_bytes().decode("utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        raise UploadValidationError(
            "json_invalid_utf8", profile.purpose, "JSON must be strict UTF-8"
        ) from exc
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise UploadValidationError(
            "json_invalid", profile.purpose, "JSON syntax is invalid"
        ) from exc
    if not isinstance(payload, dict):
        raise UploadValidationError(
            "json_schema_invalid", profile.purpose, "JSON root must be an object"
        )
    try:
        SystemSettingsRestoreDocument.model_validate_json(text, strict=True)
    except ValidationError as exc:
        raise UploadValidationError(
            "json_schema_invalid",
            profile.purpose,
            "JSON does not match the system settings restore schema",
        ) from exc


def validate_staged_file(
    *,
    purpose: UploadPurpose | str,
    path: Path,
    original_filename: str,
    declared_mime: str | None,
) -> ValidatedFile:
    profile = get_upload_profile(purpose)
    validate_upload_filename(original_filename, profile.purpose)
    size_bytes = path.stat().st_size
    if size_bytes == 0:
        raise UploadValidationError(
            "file_empty", profile.purpose, "upload must not be empty"
        )
    if size_bytes > profile.max_bytes:
        raise UploadValidationError(
            "file_too_large",
            profile.purpose,
            "upload exceeds its fixed size limit",
            size_bytes=size_bytes,
            max_bytes=profile.max_bytes,
        )
    extension, file_format = _select_format(profile, original_filename, declared_mime)
    _validate_signature(path, profile, file_format)
    if file_format.parser == "pdf":
        _validate_pdf(path, profile)
    elif file_format.parser == "image":
        _validate_image(path, profile, extension)
    elif file_format.parser == "xlsx":
        _validate_xlsx(path, profile)
    elif file_format.parser == "json":
        _validate_json(path, profile)
    else:  # pragma: no cover - immutable profiles make this unreachable
        raise RuntimeError(f"unsupported upload parser: {file_format.parser}")
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return ValidatedFile(
        purpose=profile.purpose,
        original_filename=original_filename,
        canonical_extension=extension,
        canonical_mime=file_format.canonical_mime,
        size_bytes=size_bytes,
        sha256=digest.hexdigest(),
    )
